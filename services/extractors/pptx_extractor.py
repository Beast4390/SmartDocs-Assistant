from pptx import Presentation


class PPTXExtractor:

    @staticmethod
    def extract_text(file_path):

        presentation = Presentation(file_path)

        text = []

        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text.append(shape.text)

        return "\n".join(text)